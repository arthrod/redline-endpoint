from docx import Document

# Create original document
doc1 = Document()
doc1.add_heading('Contract Agreement', 0)
doc1.add_paragraph('This agreement is made between Party A and Party B.')
doc1.add_paragraph('The term of this agreement shall be one year.')
doc1.add_paragraph('Payment shall be made monthly.')
doc1.save('/home/arthrod/workspace/redline-endpoint/original.docx')

# Create modified document
doc2 = Document()
doc2.add_heading('Contract Agreement', 0)
doc2.add_paragraph('This agreement is made between Party A and Party C.')  # Changed Party B to Party C
doc2.add_paragraph('The term of this agreement shall be two years.')  # Changed one year to two years
doc2.add_paragraph('Payment shall be made quarterly.')  # Changed monthly to quarterly
doc2.add_paragraph('Late payments will incur a 5% fee.')  # Added new paragraph
doc2.save('/home/arthrod/workspace/redline-endpoint/modified.docx')

print("Test documents created successfully!")
